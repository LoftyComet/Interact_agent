from __future__ import annotations

import hashlib

from docx import Document
from openpyxl import Workbook

from gesture_agent.indexing.extractors import extract_document
from gesture_agent.indexing.models import CorpusFile


def _source(path, role="primary") -> CorpusFile:
    content = path.read_bytes()
    return CorpusFile(
        id="file_test",
        relative_path=path.name,
        extension=path.suffix.lower(),
        media_type="application/octet-stream",
        size_bytes=len(content),
        modified_at="2026-01-01T00:00:00+00:00",
        sha256=hashlib.sha256(content).hexdigest(),
        role=role,
        authority=100,
        status="include",
    )


def test_markdown_extractor_preserves_headings_and_lines(tmp_path) -> None:
    path = tmp_path / "source.md"
    path.write_text("# 第一章\n\n第一段。\n继续。\n\n## 第二节\n内容。", encoding="utf-8")

    document = extract_document(path, _source(path))

    assert document.title == "第一章"
    assert [block.kind for block in document.blocks] == ["heading", "paragraph", "heading", "paragraph"]
    assert document.blocks[1].source_locator == {"line_start": 3, "line_end": 4}


def test_docx_extractor_preserves_heading_paragraph_and_table(tmp_path) -> None:
    path = tmp_path / "source.docx"
    source = Document()
    source.add_heading("核心机制", level=1)
    source.add_paragraph("这是定义。")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "属性"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "位置"
    table.cell(1, 1).text = "连续"
    source.save(path)

    document = extract_document(path, _source(path))

    assert document.title == "核心机制"
    assert [block.kind for block in document.blocks] == ["heading", "paragraph", "table"]
    assert "| 属性 | 值 |" in document.blocks[2].text


def test_xlsx_extractor_preserves_sheet_and_rows(tmp_path) -> None:
    path = tmp_path / "terms.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "术语"
    sheet.append(["标准术语", "别名"])
    sheet.append(["单击", "点一下"])
    workbook.save(path)

    document = extract_document(path, _source(path, role="annotation"))

    assert document.blocks[0].text == "术语"
    assert "单击" in document.blocks[1].text
