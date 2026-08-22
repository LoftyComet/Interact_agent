"""Format adapters for converting corpus files into normalized documents."""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Callable, Iterable

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import CorpusFile, DocumentBlock, NormalizedDocument


DOCUMENT_SCHEMA_VERSION = "1.0"


class ExtractionError(RuntimeError):
    """Raised when a supported source cannot be converted safely."""


def extract_document(path: str | Path, source: CorpusFile) -> NormalizedDocument:
    """Extract one source file through the adapter selected by extension."""

    source_path = Path(path)
    adapters: dict[str, Callable[[Path], tuple[list[DocumentBlock], dict]]] = {
        ".doc": _extract_legacy_doc,
        ".docx": _extract_docx,
        ".md": _extract_markdown,
        ".pdf": _extract_pdf,
        ".rtf": _extract_legacy_doc,
        ".txt": _extract_text,
        ".xls": _extract_spreadsheet,
        ".xlsx": _extract_spreadsheet,
    }
    adapter = adapters.get(source.extension)
    if adapter is None:
        raise ExtractionError(f"No extractor for {source.extension}: {source_path}")
    try:
        blocks, metadata = adapter(source_path)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Failed to extract {source_path}: {exc}") from exc
    if not any(block.text.strip() for block in blocks):
        raise ExtractionError(f"Extractor returned no text: {source_path}")
    return NormalizedDocument(
        schema_version=DOCUMENT_SCHEMA_VERSION,
        id="doc_" + source.id.removeprefix("file_"),
        file_id=source.id,
        title=_document_title(source_path, blocks),
        source_path=source.relative_path,
        source_sha256=source.sha256,
        source_format=source.extension,
        role=source.role,
        authority=source.authority,
        blocks=blocks,
        metadata=metadata,
    )


def _extract_markdown(path: Path) -> tuple[list[DocumentBlock], dict]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    blocks: list[DocumentBlock] = []
    paragraph: list[str] = []
    paragraph_start = 1

    def flush(end_line: int) -> None:
        nonlocal paragraph
        text = "\n".join(paragraph).strip()
        if text:
            blocks.append(_block(path, len(blocks), "paragraph", text, line_start=paragraph_start, line_end=end_line))
        paragraph = []

    for line_number, line in enumerate(lines, start=1):
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            flush(line_number - 1)
            blocks.append(
                _block(
                    path,
                    len(blocks),
                    "heading",
                    heading.group(2).strip(),
                    heading_level=len(heading.group(1)),
                    line_start=line_number,
                    line_end=line_number,
                )
            )
        elif not line.strip():
            flush(line_number - 1)
        else:
            if not paragraph:
                paragraph_start = line_number
            paragraph.append(line.rstrip())
    flush(len(lines))
    return blocks, {"line_count": len(lines)}


def _extract_text(path: Path) -> tuple[list[DocumentBlock], dict]:
    text = path.read_text(encoding="utf-8-sig")
    return _plain_text_blocks(path, text), {"character_count": len(text)}


def _extract_docx(path: Path) -> tuple[list[DocumentBlock], dict]:
    document = Document(path)
    blocks: list[DocumentBlock] = []
    paragraph_index = 0
    table_index = 0
    for item in _iter_docx_blocks(document):
        if isinstance(item, Paragraph):
            text = _clean_text(item.text)
            current_index = paragraph_index
            paragraph_index += 1
            if not text:
                continue
            level = _heading_level(item.style.name if item.style else "")
            kind = "heading" if level is not None else "paragraph"
            blocks.append(
                _block(
                    path,
                    len(blocks),
                    kind,
                    text,
                    heading_level=level,
                    paragraph=current_index,
                )
            )
        else:
            text = _table_to_markdown(item)
            current_index = table_index
            table_index += 1
            if text:
                blocks.append(_block(path, len(blocks), "table", text, table=current_index))
    return blocks, {
        "paragraph_count": paragraph_index,
        "table_count": table_index,
        "core_title": document.core_properties.title or "",
        "core_author": document.core_properties.author or "",
    }


def _extract_pdf(path: Path) -> tuple[list[DocumentBlock], dict]:
    reader = PdfReader(str(path))
    blocks: list[DocumentBlock] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for paragraph in _split_paragraphs(text):
            blocks.append(_block(path, len(blocks), "paragraph", paragraph, page=page_number))
    metadata = dict(reader.metadata or {})
    return blocks, {"page_count": len(reader.pages), "pdf_metadata": {str(k): str(v) for k, v in metadata.items()}}


def _extract_spreadsheet(path: Path) -> tuple[list[DocumentBlock], dict]:
    if path.suffix.lower() == ".xls":
        raise ExtractionError("Legacy .xls requires conversion to .xlsx before indexing")
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[DocumentBlock] = []
    for sheet in workbook.worksheets:
        blocks.append(_block(path, len(blocks), "heading", sheet.title, heading_level=1, sheet=sheet.title))
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                rows.append(values)
        if rows:
            blocks.append(_block(path, len(blocks), "table", _rows_to_markdown(rows), sheet=sheet.title))
    return blocks, {"sheet_count": len(workbook.sheetnames), "sheets": workbook.sheetnames}


def _extract_legacy_doc(path: Path) -> tuple[list[DocumentBlock], dict]:
    textutil = shutil.which("textutil")
    if textutil:
        result = subprocess.run(
            [textutil, "-convert", "txt", "-stdout", str(path)],
            check=False,
            capture_output=True,
        )
        if result.returncode == 0 and result.stdout:
            text = result.stdout.decode("utf-8", errors="replace")
            return _plain_text_blocks(path, text), {"converter": "textutil"}

    antiword = shutil.which("antiword")
    if antiword:
        result = subprocess.run([antiword, str(path)], check=False, capture_output=True)
        if result.returncode == 0 and result.stdout:
            text = result.stdout.decode("utf-8", errors="replace")
            return _plain_text_blocks(path, text), {"converter": "antiword"}

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory(prefix="ixdl-doc-") as temp_dir:
            result = subprocess.run(
                [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", temp_dir, str(path)],
                check=False,
                capture_output=True,
            )
            converted = Path(temp_dir) / f"{path.stem}.txt"
            if result.returncode == 0 and converted.exists():
                text = converted.read_text(encoding="utf-8", errors="replace")
                return _plain_text_blocks(path, text), {"converter": "libreoffice"}
    raise ExtractionError(f"No working legacy document converter found for {path}")


def _plain_text_blocks(path: Path, text: str) -> list[DocumentBlock]:
    return [
        _block(path, sequence, "paragraph", paragraph, paragraph=sequence)
        for sequence, paragraph in enumerate(_split_paragraphs(text))
    ]


def _split_paragraphs(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    pieces = re.split(r"\n\s*\n+", normalized)
    return [_clean_text(piece) for piece in pieces if _clean_text(piece)]


def _iter_docx_blocks(parent: DocxDocument) -> Iterable[Paragraph | Table]:
    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _heading_level(style_name: str) -> int | None:
    normalized = style_name.strip()
    if normalized.lower() == "title" or normalized in {"标题", "文档标题"}:
        return 1
    match = re.search(r"(?:heading|标题)\s*([1-6])", normalized, flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def _table_to_markdown(table: Table) -> str:
    rows = [[_clean_text(cell.text) for cell in row.cells] for row in table.rows]
    return _rows_to_markdown(rows)


def _rows_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]

    def render(row: list[str]) -> str:
        return "| " + " | ".join(value.replace("|", "\\|").replace("\n", " ") for value in row) + " |"

    return "\n".join([render(padded[0]), render(["---"] * width), *(render(row) for row in padded[1:])])


def _document_title(path: Path, blocks: list[DocumentBlock]) -> str:
    for block in blocks:
        if block.kind == "heading" and block.text:
            return block.text[:200]
    return path.stem


def _block(
    path: Path,
    sequence: int,
    kind: str,
    text: str,
    *,
    heading_level: int | None = None,
    **locator,
) -> DocumentBlock:
    stable = f"{path.name}\0{sequence}\0{kind}\0{text}".encode("utf-8")
    return DocumentBlock(
        id="block_" + hashlib.sha256(stable).hexdigest()[:20],
        sequence=sequence,
        kind=kind,  # type: ignore[arg-type]
        text=_clean_text(text),
        heading_level=heading_level,
        source_locator=locator,
    )


def _clean_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()
