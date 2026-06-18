"""把人工标注后的 Excel 回写到 data/intent_examples.json。

输入: docs/intent_examples_已标注.xlsx （由 gen_intent_xlsx.py 生成后人工标注）
输出: data/intent_examples.json

规则:
- 以「校对意图」列为准。若为空则退回「原始意图」。
- 「校对意图」可能写成非正规标签（标注者用来表达"需要新回答模版"等设计意图），
  这些不是 Intent 的合法取值，按下面规则归一到最接近的合法 intent：
    * 形如 "open_ended（新模版）" 的，取括号前的合法 intent。
    * "新模版" / "多个问题" 这类纯设计批注，退回原始意图（分类本身没变，变的是回答模版）。
    * MANUAL_OVERRIDE 里按序号给出的，优先生效（处理校对备注明确指向了别的 intent 的个例）。
- note 取「校对备注」，为空则退回「原始备注」，保留标注者的判别理由供提示词参考。
- 被归一/覆盖的条目会在控制台打印，方便人工复核。
"""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC_XLSX = ROOT / "docs" / "intent_examples_已标注.xlsx"
OUT_JSON = ROOT / "data" / "intent_examples.json"

VALID_INTENTS = {
    "basic_interaction_mechanism": "交互机制（单击/长按/拖拽/速率式等机制原理、逻辑、优缺点）",
    "control_form": "控件形态（按钮/旋钮/开关/摇杆等具体控件、控件选型、放置位置）",
    "basic_property": "基础属性（位置/角度/力/二元多级等输入属性维度）",
    "multimodal_interaction": "多模态交互（跨模态分工、非手控等）",
    "voice_interaction": "语音交互（声控/口令/对话式）",
    "interaction_compare": "交互机制对比（A 和 B 有什么区别）",
    "background_knowledge": "背景知识（什么是交互、IxDL、书的框架等）",
    "case_analysis": "理解交互案例（拆解某游戏/产品的交互、生成交互表达式）",
    "design_evaluation": "设计方案评估（给方案提建议、推荐交互方式、优化体验、评估维度）",
    "open_ended": "开放问题（与交互设计无关的闲聊/常识）",
    "mechanism_identification": "机制识别（给操作描述反推对应哪些交互机制，先给候选再追问，暂不附表达式）",
    "control_form_compare": "控件形态对比（硬件/控件载体之间的对比，非交互机制对比）",
    "function_interaction_breakdown": "功能交互拆解（拆解一类功能在多种情况下的交互枚举）",
    "mechanism_parameter_compare": "同机制参数对比（同一机制不同参数的取舍，先查论文再看模型确信度）",
    "control_form_application": "控件形态延伸应用（控件与位置/人群/场景的关系，书里没有直接内容，套话+追问+用属性间接答）",
    "interaction_optimization": "交互优化（现有交互有具体问题想优化，先诊断+追问再给建议）",
    "evaluation_methodology": "评估方法论（如何评估交互/评估维度，用控件形态/属性/机制作评估单元）",
}

# 按序号覆盖：校对备注/校对总结明确指向了与原始意图不同的合法 intent 的个例。
MANUAL_OVERRIDE = {
    8: "mechanism_identification",   # "属于什么交互机制"，反推机制而非讲解
    27: "mechanism_identification",  # "可以生成这个交互的交互表达式吗"，反推后给候选
    12: "control_form_compare",      # 手表 vs 手套，硬件载体对比
    37: "control_form_compare",      # 头部 vs 腿部控制，控件形态对比（原标 interaction_compare）
    24: "function_interaction_breakdown",  # 英雄控制涉及哪些手势，拆解一类功能
    26: "function_interaction_breakdown",  # 掉落物品各种情况的交互逻辑系统
    10: "mechanism_parameter_compare",     # 拖拽长/短距离的应用场景
    15: "control_form_application",  # 开关适合放置在什么位置
    16: "control_form_application",  # 什么样的控件形式更适合老年人
    17: "control_form_application",  # 椅子作为交互体有什么应用场景
    36: "interaction_optimization",  # 优化原神切换角色释放技能、误触
    33: "evaluation_methodology",    # 如何评估交互、评估维度有哪些
}

HEADER_NOTE = [
    "本文件由 scripts/apply_intent_xlsx.py 从 docs/intent_examples_已标注.xlsx 回写生成。",
    "intent 取人工校对后的标签；note 保留校对判别理由，用于：1) 规则识别近似命中时的强候选；2) 大模型意图判定提示词中的参考示例。",
    "可选 intent 取值见 valid_intents；若要重新标注，改 Excel 后再跑该脚本。",
]


def normalize_intent(raw_corrected: str, raw_original: str, seq: int) -> tuple[str, bool]:
    """返回 (合法 intent, 是否经过归一/覆盖)。"""
    if seq in MANUAL_OVERRIDE:
        return MANUAL_OVERRIDE[seq], True

    corrected = (raw_corrected or "").strip()
    original = (raw_original or "").strip()

    # 直接合法
    if corrected in VALID_INTENTS:
        return corrected, False

    # 形如 "open_ended（新模版）"：取括号/空白前的 token
    token = re.split(r"[（(\s]", corrected, maxsplit=1)[0].strip() if corrected else ""
    if token in VALID_INTENTS:
        return token, True

    # 纯设计批注（新模版/多个问题/空）→ 退回原始意图
    if original in VALID_INTENTS:
        return original, True

    raise ValueError(f"序号 {seq}: 无法归一意图 校对={corrected!r} 原始={original!r}")


def main() -> None:
    wb = load_workbook(SRC_XLSX, data_only=True)
    ws = wb["标注表"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}

    examples = []
    normalized = []
    for row in rows[1:]:
        if row is None or row[idx["问题"]] is None:
            continue
        seq = row[idx["序号"]]
        question = str(row[idx["问题"]]).strip()
        corrected = row[idx["校对意图"]]
        original = row[idx["原始意图"]]
        intent, changed = normalize_intent(corrected, original, seq)

        note = (row[idx["校对备注"]] or row[idx["原始备注"]] or "")
        note = str(note).strip()

        examples.append({"question": question, "intent": intent, "note": note})
        if changed:
            normalized.append((seq, str(corrected).strip(), intent, question))

    payload = {
        "// 说明": HEADER_NOTE,
        "valid_intents": VALID_INTENTS,
        "examples": examples,
    }
    OUT_JSON.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print(f"已写入: {OUT_JSON}  共 {len(examples)} 条")
    if normalized:
        print(f"\n以下 {len(normalized)} 条的标签经过归一/覆盖，请复核：")
        for seq, corrected, intent, question in normalized:
            print(f"  #{seq:>2} 校对={corrected!r:<22} -> {intent:<26} {question[:24]}")


if __name__ == "__main__":
    main()
