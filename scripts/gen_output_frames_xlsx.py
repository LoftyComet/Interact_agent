"""从 data/output_frames.json 生成方便人工标注的 Excel 表格。

输出: docs/output_frames_标注.xlsx
- 每个 intent 一行，展示当前输出框架（章节顺序，用 → 连接），预留「校对后框架」列与备注列。
- 「校对后框架」列里人工用 → 或换行分隔章节标题来调整顺序/增删章节，回写脚本再按分隔符解析。
- 另附「框架明细」表，把每个 intent 的章节逐条拆开（序号 + 章节名），方便逐节标注。
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from gesture_agent.learning.question_parser import INTENT_LABELS

SRC = ROOT / "data" / "output_frames.json"
OUT = ROOT / "docs" / "output_frames_标注.xlsx"

SEP = " → "

data = json.loads(SRC.read_text(encoding="utf-8"))
frames = data["frames"]

wb = Workbook()

header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(color="FFFFFF", bold=True)
wrap = Alignment(wrap_text=True, vertical="top")


def style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")


# ---- Sheet1: 框架标注表（每个 intent 一行）----
ws = wb.active
ws.title = "框架标注表"
headers = ["序号", "意图键", "意图中文名", "当前输出框架", "校对后框架", "备注"]
ws.append(headers)
style_header(ws, len(headers))

for i, (intent, sections) in enumerate(frames.items(), start=1):
    current = SEP.join(sections)
    ws.append([
        i,
        intent,
        INTENT_LABELS.get(intent, ""),
        current,
        current,   # 校对后框架默认填当前框架，人工在此调整
        "",
    ])

last_row = len(frames) + 1
widths = [6, 32, 18, 60, 60, 30]
for col, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(col)].width = w
for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=len(headers)):
    for c in row:
        c.alignment = wrap
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"

# ---- Sheet2: 框架明细（每个章节一行）----
ws2 = wb.create_sheet("框架明细")
headers2 = ["意图键", "意图中文名", "章节序号", "章节标题", "校对后章节标题", "是否保留", "备注"]
ws2.append(headers2)
style_header(ws2, len(headers2))

for intent, sections in frames.items():
    for idx, title in enumerate(sections, start=1):
        ws2.append([
            intent,
            INTENT_LABELS.get(intent, ""),
            idx,
            title,
            title,   # 校对后章节标题默认填原标题
            "是",     # 是否保留：是/否
            "",
        ])

detail_rows = sum(len(v) for v in frames.values()) + 1
widths2 = [32, 18, 8, 24, 24, 10, 30]
for col, w in enumerate(widths2, start=1):
    ws2.column_dimensions[get_column_letter(col)].width = w
for row in ws2.iter_rows(min_row=2, max_row=detail_rows, min_col=1, max_col=len(headers2)):
    for c in row:
        c.alignment = wrap
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{detail_rows}"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"已生成: {OUT}  共 {len(frames)} 个意图，{detail_rows - 1} 个章节")
