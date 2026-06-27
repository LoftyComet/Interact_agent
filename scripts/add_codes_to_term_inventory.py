"""
从 Excel 解析路径代号，更新 term_inventory.json 树节点的 code 字段。

用法:
    uv run python scripts/add_codes_to_term_inventory.py
"""

import json
import re
import openpyxl


def clean_mechanism_label(text: str) -> str:
    """
    从冲突调和/多维协同的 Excel 文本中提取术语名。

    Excel 格式如:
      "单击 🆚 长按：快击/快抬（010变化/时间）" → "快击/快抬"
      "同向拖拽：边缘滑入（起始选择机制）" → "边缘滑入"
      "一次性 🆚 持续性 点击：缓冲连发" → "缓冲连发"
    """
    if not text:
        return None
    text = str(text).strip()
    # Extract text after "：" (fullwidth colon)
    if "：" in text:
        text = text.split("：")[-1].strip()
    # Extract text after ":" (halfwidth colon) but only if not part of time
    elif ":" in text and not text.startswith("例"):
        parts = text.split(":")
        if len(parts) > 1:
            text = parts[-1].strip()
    # Remove trailing parenthetical like "（010变化/时间）"
    text = re.sub(r'[（(][^)）]*[)）]\s*$', '', text).strip()
    # Remove content after 🆚 (for remaining cases)
    if '🆚' in text:
        text = text.split('🆚')[0].strip()
    # Remove newlines and subsequent examples
    text = text.split('\n')[0].strip()
    return text


def clean_simple_label(text: str) -> str:
    """简单清洗：去括号、去换行"""
    if not text:
        return None
    text = str(text).strip()
    text = re.sub(r'[（(][^)）]*[)）]', '', text).strip()
    text = text.split('\n')[0].strip()
    return text


def build_excel_code_map(ws) -> dict:
    """
    从 Excel 构建 section → [(code, label), ...] 的映射。
    """
    sections = {}

    # === 基础属性 (Col B=code, Col C=name) rows 3-14 ===
    items = []
    for r in range(3, 15):
        code = ws.cell(row=r, column=2).value
        name = ws.cell(row=r, column=3).value
        if code and name:
            try:
                code_str = str(int(code))
            except (ValueError, TypeError):
                continue
            name = clean_simple_label(name)
            if name:
                items.append((code_str, name))
    sections['基础属性'] = items

    # === 点击类 (Col D=code, Col E=name) rows 4-14 ===
    items = []
    for r in range(4, 15):
        code = ws.cell(row=r, column=4).value
        name = ws.cell(row=r, column=5).value
        if code and name and '-' in str(code):
            code_str = str(code).strip()
            name = clean_simple_label(name)
            if name:
                items.append((code_str, name))
    sections['点击类'] = items

    # === 位移类 (Col F=code, Col G=name) rows 4-22 ===
    items = []
    for r in range(4, 23):
        code = ws.cell(row=r, column=6).value
        name = ws.cell(row=r, column=7).value
        if code and name and '-' in str(code):
            code_str = str(code).strip()
            name = clean_simple_label(name)
            if name:
                items.append((code_str, name))
    sections['位移类'] = items

    # === 多维协同 (Col H=code, Col I=name) rows 3-18 ===
    items = []
    for r in range(3, 19):
        code = ws.cell(row=r, column=8).value
        name = ws.cell(row=r, column=9).value
        if code and name and '-' in str(code):
            code_str = str(code).strip()
            name = clean_mechanism_label(name) or clean_simple_label(name)
            if name:
                items.append((code_str, name))
    sections['多维协同'] = items

    # === 冲突调和 (Col J=code, Col K=name) rows 4-14 ===
    items = []
    for r in range(4, 15):
        code = ws.cell(row=r, column=10).value
        name = ws.cell(row=r, column=11).value
        if code and name and '-' in str(code):
            code_str = str(code).strip()
            name = clean_mechanism_label(name) or clean_simple_label(name)
            if name:
                items.append((code_str, name))
    sections['冲突调和'] = items

    # === 实体 Object (Col L=code, Col M=name) rows 5-30 ===
    items = []
    for r in range(5, 31):
        code = ws.cell(row=r, column=12).value
        name = ws.cell(row=r, column=13).value
        if code and name:
            try:
                code_str = str(int(code))
            except (ValueError, TypeError):
                continue
            name = clean_simple_label(name)
            # Skip section headers like 【点触实体】
            if name and not name.startswith('【'):
                items.append((code_str, name))
    sections['实体'] = items

    # === 姿态 Gesture (Col N=code, Col O=name) rows 5-17 ===
    items = []
    for r in range(5, 18):
        code = ws.cell(row=r, column=14).value
        name = ws.cell(row=r, column=15).value
        if code and name:
            try:
                code_str = str(int(code))
            except (ValueError, TypeError):
                continue
            name = clean_simple_label(name)
            if name and not name.startswith('【'):
                items.append((code_str, name))
    sections['姿态'] = items

    # === 流体 Fluid (Col N=code, Col O=name) rows 20-22 ===
    items = []
    for r in range(20, 23):
        code = ws.cell(row=r, column=14).value
        name = ws.cell(row=r, column=15).value
        if code and name:
            try:
                code_str = str(int(code))
            except (ValueError, TypeError):
                continue
            name = clean_simple_label(name)
            if name and not name.startswith('【'):
                items.append((code_str, name))
    sections['流体'] = items

    # === 其他 Others (Col N=code, Col O=name) rows 25-27 ===
    items = []
    for r in range(25, 28):
        code = ws.cell(row=r, column=14).value
        name = ws.cell(row=r, column=15).value
        if code and name:
            try:
                code_str = str(int(code))
            except (ValueError, TypeError):
                continue
            name = clean_simple_label(name)
            if name and not name.startswith('【'):
                items.append((code_str, name))
    sections['其他'] = items

    # === 多模态组合案例 (Col P=code, Col Q=name) rows 26-30 ===
    items = []
    for r in range(26, 31):
        code = ws.cell(row=r, column=16).value
        name = ws.cell(row=r, column=17).value
        if code and name:
            try:
                code_str = str(int(code))
            except (ValueError, TypeError):
                continue
            name = clean_simple_label(name)
            if name:
                items.append((code_str, name))
    sections['多模态组合案例'] = items

    return sections


def normalize(s: str) -> str:
    """标准化字符串用于比对"""
    s = s.replace('·', '/').replace(' ', '').replace('+', '+')
    return s


def match_label(json_label: str, excel_label: str, json_node: dict = None) -> bool:
    """
    判断 JSON 节点 label 是否匹配 Excel label。
    支持多种匹配策略。
    """
    # 1. 精确匹配
    if json_label == excel_label:
        return True

    # 2. 标准化后匹配（处理 · vs / 的差异）
    if normalize(json_label) == normalize(excel_label):
        return True

    # 3. JSON label 以 Excel label 开头（如 "二元属性" 匹配 "二元"）
    if json_label.startswith(excel_label) and len(excel_label) >= 2:
        return True

    # 4. Excel label 以 JSON label 开头
    if excel_label.startswith(json_label) and len(json_label) >= 2:
        return True

    # 5. 通过 alias 匹配
    if json_node:
        aliases = json_node.get('aliases', [])
        if isinstance(aliases, list):
            for alias in aliases:
                if normalize(alias) == normalize(excel_label) or alias == excel_label:
                    return True

    return False


def apply_codes(tree_node: dict, section_codes: list, section_label: str, depth: int = 0) -> dict:
    """
    递归遍历 JSON 树，对匹配 section 的节点子树应用 code。

    匹配逻辑:
    - 如果 section_codes 不为空且 tree_node.label 匹配 section_label:
      对 tree_node.children 中的每个子节点，查找 section_codes 中匹配的 code
    - 否则继续递归查找
    """
    if not tree_node:
        return tree_node

    node_label = tree_node.get('label', '')
    children = tree_node.get('children', [])

    # Check if this node is the target section
    if section_codes and node_label == section_label and children:
        # Match each child to a code
        for child in children:
            child_label = child.get('label', '')
            for code, excel_label in section_codes:
                if match_label(child_label, excel_label, child):
                    child['code'] = code
                    break
        return tree_node

    # Recurse into children
    if children:
        for child in children:
            apply_codes(child, section_codes, section_label, depth + 1)

    return tree_node


def main():
    # 1. 加载 Excel
    wb = openpyxl.load_workbook('data/手势词典大纲定稿.xlsx')
    ws = wb.active
    excel_sections = build_excel_code_map(ws)

    print("=== Excel sections parsed ===")
    for sec, items in excel_sections.items():
        print(f"  {sec}: {len(items)} items")
        for code, name in items:
            print(f"    {code}: {name}")

    # 2. 加载现有 JSON
    with open('data/term_inventory.json', 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 3. 定义 section 到 JSON 树路径的映射
    # 每个条目: (section_name, target_node_label, excel_section_key)
    section_targets = [
        ('基础属性', '基础属性', '基础属性'),
        ('交互逻辑', '点击类', '点击类'),
        ('交互逻辑', '位移类', '位移类'),
        ('交互逻辑', '多维协同', '多维协同'),
        ('交互逻辑', '冲突调和', '冲突调和'),
        ('控件形态', '实体', '实体'),
        ('控件形态', '姿态', '姿态'),
        ('控件形态', '流体', '流体'),
        ('控件形态', '其他', '其他'),
        ('含义与多模态', '多模态组合案例', '多模态组合案例'),
    ]

    # 4. 为每个 section 应用 codes
    root = data['root']

    for parent_label, target_label, excel_key in section_targets:
        codes = excel_sections.get(excel_key, [])
        if not codes:
            print(f"  WARNING: No codes found for {excel_key}")
            continue

        # Find parent node
        parent = None
        for child in root.get('children', []):
            if child.get('label') == parent_label:
                parent = child
                break

        if parent:
            # Find and process target node
            for subchild in parent.get('children', []):
                if subchild.get('label') == target_label:
                    targets = subchild.get('children', [])
                    matched = 0
                    for t in targets:
                        t_label = t.get('label', '')
                        for code, excel_label in codes:
                            if match_label(t_label, excel_label, t):
                                t['code'] = code
                                matched += 1
                                break
                    print(f"  {excel_key}: matched {matched}/{len(targets)} children")
                    break
        else:
            print(f"  WARNING: Parent '{parent_label}' not found")

    # 5. Also add codes to section-level group nodes
    # "点击类" → code "1", "位移类" → code "2", etc.
    group_codes = {
        '点击类': '1',
        '位移类': '2',
        '多维协同': '3',
        '冲突调和': '4',
    }
    for child in root.get('children', []):
        if child.get('label') == '交互逻辑':
            for sub in child.get('children', []):
                if sub.get('label') in group_codes:
                    sub['code'] = group_codes[sub['label']]

    # 6. 保存
    with open('data/term_inventory.json', 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print("\n✅ term_inventory.json updated with code fields")


if __name__ == '__main__':
    main()
