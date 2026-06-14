"""从 data/intent_examples.json 生成方便人工标注的 Excel 表格。

输出: docs/intent_examples_标注.xlsx
- 每条问题一行，保留原始意图与含义，预留「校对意图」下拉列与备注列。
- 「校对意图」列使用数据校验下拉框，取值来自 valid_intents。
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "intent_examples.json"
OUT = ROOT / "docs" / "intent_examples_标注.xlsx"

data = json.loads(SRC.read_text(encoding="utf-8"))
valid = data["valid_intents"]
examples = data["examples"]

wb = Workbook()

# ---- Sheet1: 标注表 ----
ws = wb.active
ws.title = "标注表"

headers = ["序号", "问题", "原始意图", "原始意图含义", "校对意图", "校对备注", "原始备注"]
ws.append(headers)

header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(color="FFFFFF", bold=True)
for col, _ in enumerate(headers, start=1):
    c = ws.cell(row=1, column=col)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for i, ex in enumerate(examples, start=1):
    intent = ex.get("intent", "")
    ws.append([
        i,
        ex.get("question", ""),
        intent,
        valid.get(intent, ""),
        intent,            # 校对意图默认填原始意图
        "",                # 校对备注留空
        ex.get("note", ""),
    ])

# 下拉校验：校对意图列(第5列)
intent_keys = list(valid.keys())
formula = '"' + ",".join(intent_keys) + '"'
dv = DataValidation(type="list", formula1=formula, allow_blank=True)
dv.error = "请从下拉列表中选择有效的意图"
dv.errorTitle = "无效意图"
dv.prompt = "从下拉框选择校对后的意图"
dv.promptTitle = "校对意图"
ws.add_data_validation(dv)
last_row = len(examples) + 1
dv.add(f"E2:E{last_row}")

# 列宽
widths = [6, 60, 26, 50, 26, 24, 24]
for col, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(col)].width = w

# 自动换行 + 顶端对齐
wrap = Alignment(wrap_text=True, vertical="top")
for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=len(headers)):
    for c in row:
        c.alignment = wrap

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"

# ---- Sheet2: 意图说明 ----
ws2 = wb.create_sheet("意图说明")
ws2.append(["意图键", "含义"])
for col in (1, 2):
    c = ws2.cell(row=1, column=col)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")
for k, v in valid.items():
    ws2.append([k, v])
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 70
for row in ws2.iter_rows(min_row=2, max_row=len(valid) + 1, min_col=1, max_col=2):
    for c in row:
        c.alignment = wrap
ws2.freeze_panes = "A2"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"已生成: {OUT}  共 {len(examples)} 条")
