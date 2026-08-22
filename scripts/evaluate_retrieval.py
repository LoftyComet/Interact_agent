#!/usr/bin/env python3
"""Evaluate local intent parsing and retrieval against an annotated XLSX set."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from gesture_agent.knowledge import KnowledgeBase
from gesture_agent.learning import QuestionParser


LEGACY_INTENT_MAP = {
    "control_form_compare": "interaction_compare",
    "control_form_application": "design_suggestion",
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("dataset", type=Path)
    parser.add_argument("--index-dir", type=Path, default=Path("knowledge_index"))
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--output", type=Path, default=Path("knowledge_index/evaluation_baseline.json"))
    parser.add_argument("--top-k", type=int, default=6)
    args = parser.parse_args()

    kb = KnowledgeBase.load(args.data_dir, index_dir=args.index_dir)
    question_parser = QuestionParser(kb)
    terminology = json.loads((args.index_dir / "terminology.json").read_text(encoding="utf-8"))
    terms = sorted(
        (str(item["label"]) for item in terminology.get("terms", []) if len(str(item["label"])) >= 2),
        key=lambda value: (-len(value), value),
    )
    rows = _load_rows(args.dataset)
    details: list[dict] = []
    raw_intent_hits = 0
    intent_hits = 0
    intent_query_count = 0
    nonempty_hits = 0
    primary_hits = 0
    queries_with_terms = 0
    queries_with_term_hit = 0
    term_recalls: list[float] = []

    for row_number, prompt, expected_output, expected_intent in rows:
        structure = question_parser.parse(prompt)
        normalized_expected_intent = _normalize_expected_intent(expected_intent)
        hits = kb.retriever.retrieve(
            prompt,
            top_k=args.top_k,
            preferred_terms=structure.terms,
        )
        expected_terms = [term for term in terms if term in expected_output]
        retrieved_text = "\n".join(
            f"{hit.title}\n{hit.context_text}\n{' '.join(hit.terms)}" for hit in hits
        )
        matched_terms = [term for term in expected_terms if term in retrieved_text]
        if structure.intent == expected_intent:
            raw_intent_hits += 1
        if normalized_expected_intent is not None:
            intent_query_count += 1
        if structure.intent == normalized_expected_intent:
            intent_hits += 1
        if hits:
            nonempty_hits += 1
        if any(hit.authority >= 100 for hit in hits):
            primary_hits += 1
        if expected_terms:
            queries_with_terms += 1
            recall = len(matched_terms) / len(expected_terms)
            term_recalls.append(recall)
            if matched_terms:
                queries_with_term_hit += 1
        details.append({
            "row": row_number,
            "query": prompt,
            "dataset_intent": expected_intent,
            "expected_intent": normalized_expected_intent,
            "intent_evaluable": normalized_expected_intent is not None,
            "predicted_intent": structure.intent,
            "parsed_terms": structure.terms,
            "expected_terms": expected_terms,
            "matched_expected_terms": matched_terms,
            "hits": [
                {
                    "chunk_id": hit.chunk_id,
                    "title": hit.title,
                    "source": hit.source_path,
                    "authority": hit.authority,
                    "score": hit.score,
                }
                for hit in hits
            ],
        })

    count = len(rows)
    metrics = {
        "query_count": count,
        "top_k": args.top_k,
        "retrieval_mode": kb.retriever.mode if kb.retriever else "legacy",
        "intent_accuracy": _ratio(intent_hits, intent_query_count),
        "intent_query_count": intent_query_count,
        "raw_dataset_intent_accuracy": _ratio(raw_intent_hits, count),
        "excluded_feedback_intent_count": count - intent_query_count,
        "nonempty_retrieval_rate": _ratio(nonempty_hits, count),
        "primary_source_hit_rate": _ratio(primary_hits, count),
        "expected_term_query_count": queries_with_terms,
        "expected_term_hit_rate": _ratio(queries_with_term_hit, queries_with_terms),
        "macro_expected_term_recall": round(sum(term_recalls) / len(term_recalls), 4) if term_recalls else 0.0,
    }
    payload = {
        "schema_version": "1.0",
        "dataset": str(args.dataset.resolve()),
        "index_build_id": json.loads((args.index_dir / "manifest.json").read_text(encoding="utf-8"))["build_id"],
        "metrics": metrics,
        "queries": details,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output.resolve()), **metrics}, ensure_ascii=False, indent=2))


def _load_rows(path: Path) -> list[tuple[int, str, str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows: list[tuple[int, str, str, str]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        prompt = str(values[1] or "").strip()
        if not prompt:
            continue
        rows.append((row_number, prompt, str(values[2] or ""), str(values[3] or "").strip()))
    return rows


def _ratio(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 4) if denominator else 0.0


def _normalize_expected_intent(raw_intent: str) -> str | None:
    if raw_intent.startswith("feedback_"):
        # These labels describe a multi-turn product feedback taxonomy, not the
        # single-turn domain intent contract used by QuestionParser.
        return None
    return LEGACY_INTENT_MAP.get(raw_intent, raw_intent)


if __name__ == "__main__":
    main()
